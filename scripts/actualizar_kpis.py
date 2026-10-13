#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
actualizar_kpis.py
==================
Carga el "Formato de Seguimiento de KPI MTDH" (Excel con las hojas Catálogo_KPIs y
Seguimiento_Mensual) hacia los CSV que lee el portal.

    python scripts/actualizar_kpis.py --archivo "Seguimiento_KPI_septiembre_2026.xlsx" --area pueblos

    # Indicadores en porcentaje que llegan como hoja de cálculo del mes (numerador / denominador),
    # p. ej. KPI_Septiembre_2026_PRPJRD.xlsx -> se indica a qué KPI del catálogo corresponde:
    python scripts/actualizar_kpis.py --archivo "KPI_Septiembre_2026_PRPJRD.xlsx" --area pueblos --kpi KPI-2026-01

Qué hace:
  1. Lee el catálogo (una fila por KPI) y el seguimiento (una fila por KPI y mes).
  2. NO usa las fórmulas del Excel: recalcula % mensual, acumulados y % de avance a partir
     de las celdas que se digitan (meta programada y meta alcanzada del mes). Así no le
     afectan la clave auxiliar TEXT(...,"YYYYMM") ni los rangos fijos del Resumen.
  3. Agrega o reemplaza los meses del archivo sin duplicar: funciona igual si el Excel trae
     todo el histórico o solo el mes nuevo. Si un mes viene repetido, gana la versión más alta.
  3b. Conserva los KPI del catálogo publicado que no vengan en el archivo (no los borra).
  3c. KPI con unidad «Porcentaje»: la meta programada es el denominador y la meta alcanzada el numerador
      del mes; el resultado del mes es numerador ÷ denominador × 100 y el acumulado del año es
      Σ numeradores ÷ Σ denominadores × 100 (no se usa meta anual).
  4. No publica nombres de personas (de "Elaborado por" / "Revisado por" conserva solo el cargo)
     ni rutas de red (solo enlaces http/https).
  5. Escribe:
       datos/kpis/<area>/catalogo.csv
       datos/kpis/<area>/seguimiento.csv          (histórico, una fila por KPI × mes)
       datos/kpis/kpi_pueblos_nacionalidades.csv        (tarjetas del "Panorama destacado" de kpis.html)
       reportes/kpis/validacion_<area>_<AAAA-MM>.md
"""
import argparse
import csv
import datetime as dt
import os
import re
import sys
import warnings
from collections import defaultdict

warnings.filterwarnings("ignore")
import openpyxl

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]
MES_ABR = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
ESTADOS = ["Cumplido", "En proceso", "Con alerta", "Incumplido", "No aplica"]
AREAS = {  # área del portal -> archivo del panorama de kpis.html y ruta de la sección
    "pueblos": {"panorama": "kpi_pueblos_nacionalidades.csv", "area_kpis": "kpis-pueblos", "nombre": "Pueblos y Nacionalidades"},
}
CAMPOS_CAT = ["id_kpi", "institucion", "tipo", "resultado_esperado", "indicador", "periodicidad", "unidad",
              "linea_base", "meta_anual", "anio_meta", "responsable", "observaciones",
              "definicion", "formula", "fuente", "desagregacion", "disponibilidad", "etiqueta_numerador", "etiqueta_denominador"]
CAMPOS_SEG = ["id_kpi", "periodo", "meta_programada", "meta_alcanzada", "cumplimiento_mensual_pct",
              "programado_acumulado", "avance_acumulado", "avance_anual_pct", "cumplimiento_fecha_pct",
              "estado", "logros", "alertas", "medios_verificacion", "enlace", "detalle_casos",
              "elaborado_cargo", "revisado_cargo", "fecha_elaboracion", "version", "fecha_actualizacion", "validacion"]


def num(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("%", "")
    if re.fullmatch(r"\d{1,3}(\.\d{3})+(,\d+)?", s):
        s = s.replace(".", "").replace(",", ".")
    s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def periodo(v):
    """Fecha de Excel, '2026-08', '08/2026' o 'agosto 2026' -> 'AAAA-MM'."""
    if isinstance(v, (dt.datetime, dt.date)):
        return f"{v.year}-{v.month:02d}"
    s = str(v or "").strip().lower()
    m = re.match(r"(\d{4})[-/](\d{1,2})", s) or None
    if m:
        return f"{m.group(1)}-{int(m.group(2)):02d}"
    m = re.match(r"(\d{1,2})[-/](\d{4})", s)
    if m:
        return f"{m.group(2)}-{int(m.group(1)):02d}"
    for i, nombre in enumerate(MESES):
        if nombre in s:
            a = re.search(r"\d{4}", s)
            if a:
                return f"{a.group(0)}-{i + 1:02d}"
    return ""


def fecha(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    return str(v or "").strip()[:10]


def solo_cargo(v):
    """'Nombre Apellido — Cargo' -> 'Cargo' (no se publican nombres de personas).
    Si la celda trae varias personas (una por línea), conserva solo el cargo de cada una."""
    cargos = []
    for linea in str(v or "").splitlines():
        partes = re.split(r"\s+[—–-]\s+", linea.strip(), maxsplit=1)
        if len(partes) == 2 and partes[1].strip():
            cargos.append(partes[1].strip())
    return " / ".join(cargos)


def fnum(x, d=2):
    if x is None or x == "":
        return ""
    x = round(float(x), d)
    return str(int(x)) if x == int(x) else str(x)


def per_txt(p):
    return f"{MESES[int(p[5:7]) - 1]} {p[:4]}"


def leer_csv(ruta):
    if not os.path.exists(ruta):
        return []
    with open(ruta, encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def escribir_csv(ruta, filas, campos):
    os.makedirs(os.path.dirname(ruta), exist_ok=True)
    with open(ruta, "w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=campos, extrasaction="ignore")
        w.writeheader()
        w.writerows(filas)


def encabezados(ws, fila=2):
    """{texto normalizado del encabezado: columna} para no depender de la posición exacta."""
    out = {}
    for c in ws[fila]:
        if c.value:
            out[re.sub(r"\s+", " ", str(c.value)).strip().lower()] = c.column
    return out


def col(h, *claves):
    for k in claves:
        for texto, c in h.items():
            if texto.startswith(k):
                return c
    return None


def es_pct(c):
    return bool(re.search(r"porcentaje|%", str(c.get("unidad") or ""), re.I))


def leer_calculo(ruta, id_kpi):
    """Hoja «CÁLCULO DEL INDICADOR — <MES> <AÑO>»: filas etiqueta | valor con Periodo, Numerador, Denominador
    e Interpretación. Devuelve una fila de seguimiento (numerador -> meta alcanzada, denominador -> programada)."""
    wb = openpyxl.load_workbook(ruta, data_only=True)
    pares, textos = {}, []
    for ws in wb.worksheets:
        for fila in ws.iter_rows(values_only=True):
            celdas = [c for c in fila if c not in (None, "")]
            if not celdas:
                continue
            etq = str(celdas[0]).strip().lower()
            if len(celdas) >= 2:
                pares.setdefault(etq, celdas[1])
            elif etq.startswith("interpretaci"):
                textos.append(str(celdas[0]).strip())
    def buscar(*claves):
        for k, v in pares.items():
            if any(k.startswith(c) for c in claves):
                return v
        return None
    interp = next((str(v) for k, v in pares.items() if k.startswith("interpretaci")), None) or (textos[0] if textos else "")
    interp = re.sub(r"^interpretaci[oó]n\s*:\s*", "", interp, flags=re.I).strip()
    p = periodo(buscar("periodo", "período"))
    nu, de = num(buscar("numerador")), num(buscar("denominador"))
    if interp and nu is not None and de:
        # El porcentaje del texto viene de TEXT(...;"0.00"), que en Excel en español sale mal (p. ej. «070 %»):
        # se reemplaza por el valor recalculado.
        interp = re.sub(r"\bel\s+[\d.,]+\s*%", "el " + f"{nu / de * 100:.1f}".replace(".", ",") + " %", interp, count=1)
        interp = interp[:1].upper() + interp[1:]
    return [{"fila": 0, "id_kpi": id_kpi, "periodo": p, "periodo_raw": buscar("periodo", "período"),
             "meta_programada": de, "meta_alcanzada": nu,
             "estado": "", "logros": interp, "alertas": "", "medios_verificacion": "", "enlace": "", "detalle_casos": "",
             "elaborado": "", "revisado": "", "fecha_elaboracion": "", "version": 1, "fecha_actualizacion": ""}]


def leer_excel(ruta):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    C, S = wb["Catálogo_KPIs"], wb["Seguimiento_Mensual"]
    hc = encabezados(C)
    anio_meta = ""
    for texto in hc:
        m = re.search(r"meta anual (\d{4})", texto)
        if m:
            anio_meta = m.group(1)
    cat = []
    for r in range(3, C.max_row + 1):
        g = lambda *k: C.cell(r, col(hc, *k)).value if col(hc, *k) else None
        idk = str(g("id_kpi") or "").strip()
        if not idk:
            continue
        cat.append({"id_kpi": idk, "institucion": g("institución", "institucion"), "tipo": g("tipo de indicador"),
                    "resultado_esperado": g("kpi estratégico", "kpi estrategico"), "indicador": g("indicador"),
                    "periodicidad": g("periodicidad"), "unidad": g("unidad de medida"), "linea_base": g("línea base", "linea base"),
                    "meta_anual": num(g("meta anual")), "anio_meta": anio_meta, "responsable": g("responsable"),
                    "observaciones": g("observaciones")})
    hs = encabezados(S)
    seg = []
    for r in range(3, S.max_row + 1):
        g = lambda *k: S.cell(r, col(hs, *k)).value if col(hs, *k) else None
        idk = str(g("id_kpi") or "").strip()
        if not idk:
            continue
        seg.append({"fila": r, "id_kpi": idk, "periodo": periodo(g("período de reporte", "periodo de reporte")),
                    "periodo_raw": g("período de reporte", "periodo de reporte"),
                    "meta_programada": num(g("meta programada")), "meta_alcanzada": num(g("meta alcanzada")),
                    "estado": str(g("estado del indicador") or "").strip(), "logros": str(g("logros") or "").strip(),
                    "alertas": str(g("retrasos") or "").strip(), "medios_verificacion": str(g("medios de verificación", "medios de verificacion") or "").strip(),
                    "enlace": str(g("enlace") or "").strip(), "detalle_casos": str(g("detalle de casos") or "").strip(),
                    "elaborado": g("elaborado por"), "revisado": g("revisado"), "fecha_elaboracion": fecha(g("fecha de elaboración", "fecha de elaboracion")),
                    "version": num(g("n° de versión", "nº de versión", "n° de version")) or 1, "fecha_actualizacion": fecha(g("fecha de última", "fecha de ultima"))})
    return cat, seg


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--archivo", required=True, help="Excel del formato de seguimiento de KPI")
    ap.add_argument("--area", default="pueblos", choices=list(AREAS))
    ap.add_argument("--datos", default=os.path.join(REPO, "datos"))
    ap.add_argument("--reportes", default=os.path.join(REPO, "reportes", "kpis"))
    ap.add_argument("--kpi", help="ID del KPI del catálogo publicado al que corresponde una hoja de cálculo mensual "
                                   "(numerador / denominador) que no trae el formato de seguimiento")
    args = ap.parse_args()
    cfg = AREAS[args.area]
    dir_area = os.path.join(args.datos, "kpis", args.area)

    hall = []
    H = lambda sev, cod, donde, txt: hall.append((sev, cod, donde, txt))
    cat_pub = leer_csv(os.path.join(dir_area, "catalogo.csv"))
    if args.kpi:
        if args.kpi not in {c["id_kpi"] for c in cat_pub}:
            sys.exit(f"{args.kpi} no existe en datos/kpis/{args.area}/catalogo.csv: agréguelo primero al catálogo.")
        cat, seg_xls = [], leer_calculo(args.archivo, args.kpi)
    else:
        cat, seg_xls = leer_excel(args.archivo)
    # Catálogo = el publicado, actualizado con lo que trae el Excel (los KPI que no vienen se conservan)
    ids = {c["id_kpi"]: {**c, "meta_anual": num(c.get("meta_anual"))} for c in cat_pub}
    for c in cat:
        ids[c["id_kpi"]] = {**ids.get(c["id_kpi"], {}), **{k: v for k, v in c.items() if v not in (None, "")}}
    ids = dict(sorted(ids.items()))          # orden estable por código de KPI
    cat = list(ids.values())
    if not cat:
        sys.exit("El catálogo está vacío.")

    # ---- validación de filas del Excel
    por_clave = {}
    for s in seg_xls:
        donde = f"{s['id_kpi']} · fila {s['fila']}"
        if s["id_kpi"] not in ids:
            H("ERROR", "K1", donde, "ID_KPI que no existe en el catálogo: la fila no se carga.")
            continue
        if not s["periodo"]:
            H("ERROR", "K2", donde, f"Período de reporte no reconocible: «{s['periodo_raw']}».")
            continue
        if s["meta_programada"] is None or s["meta_alcanzada"] is None:
            H("ERROR", "K3", donde, "Falta la meta programada o la meta alcanzada del mes (celdas amarillas).")
            continue
        clave = (s["id_kpi"], s["periodo"])
        if clave in por_clave:
            previa = por_clave[clave]
            H("ADVERTENCIA", "K4", donde, f"El mes {s['periodo']} está repetido (filas {previa['fila']} y {s['fila']}); se usa la versión más alta "
              "y no se suma dos veces.")
            if s["version"] < previa["version"]:
                continue
        por_clave[clave] = s

    # ---- fusión con el histórico ya publicado (el Excel puede traer todo o solo el mes nuevo)
    hist = {(r["id_kpi"], r["periodo"]): r for r in leer_csv(os.path.join(dir_area, "seguimiento.csv"))}
    nuevos, reemplazados = 0, 0
    for clave, s in por_clave.items():
        if clave in hist:
            reemplazados += 1
        else:
            nuevos += 1
        hist[clave] = {
            "id_kpi": s["id_kpi"], "periodo": s["periodo"], "meta_programada": fnum(s["meta_programada"]),
            "meta_alcanzada": fnum(s["meta_alcanzada"]), "estado": s["estado"], "logros": s["logros"], "alertas": s["alertas"],
            "medios_verificacion": s["medios_verificacion"], "enlace": s["enlace"] if re.match(r"https?://", s["enlace"]) else "",
            "detalle_casos": s["detalle_casos"], "elaborado_cargo": solo_cargo(s["elaborado"]), "revisado_cargo": solo_cargo(s["revisado"]),
            "fecha_elaboracion": s["fecha_elaboracion"], "version": fnum(s["version"]), "fecha_actualizacion": s["fecha_actualizacion"]}
        if s["enlace"] and not re.match(r"https?://", s["enlace"]):
            H("INFO", "K9", f"{s['id_kpi']} · {s['periodo']}", f"El enlace del soporte no es una dirección web («{s['enlace'][:60]}»): no se publica.")

    # ---- recálculo por KPI, en orden cronológico
    filas = []
    ultimo = {}
    for idk, c in ids.items():
        rs = sorted([r for (i, _), r in hist.items() if i == idk], key=lambda r: r["periodo"])
        if not rs:
            H("ADVERTENCIA", "K5", idk, "El KPI no tiene ningún mes reportado.")
            continue
        meta = c["meta_anual"]
        acum = prog = 0.0
        anio_prev = None
        for r in rs:
            anio = r["periodo"][:4]
            if anio != anio_prev:              # el acumulado se reinicia cada año (meta anual)
                acum = prog = 0.0
                anio_prev = anio
            p, a = num(r["meta_programada"]) or 0, num(r["meta_alcanzada"]) or 0
            acum += a
            prog += p
            r["cumplimiento_mensual_pct"] = fnum(a / p * 100) if p else ""
            r["programado_acumulado"] = fnum(prog)
            r["avance_acumulado"] = fnum(acum)
            r["avance_anual_pct"] = fnum(acum / meta * 100) if (meta and not es_pct(c)) else ""
            r["cumplimiento_fecha_pct"] = fnum(acum / prog * 100) if prog else ""
            cods = []
            if es_pct(c) and a > p:
                H("ERROR", "K12", f"{idk} · {r['periodo']}", f"El numerador ({fnum(a)}) es mayor que el denominador ({fnum(p)}).")
                cods.append("K12")
            if r["estado"] and r["estado"] not in ESTADOS:
                H("ADVERTENCIA", "K6", f"{idk} · {r['periodo']}", f"Estado «{r['estado']}» fuera de la lista ({', '.join(ESTADOS)}).")
                cods.append("K6")
            if meta and acum > meta and not es_pct(c):
                H("ADVERTENCIA", "K7", f"{idk} · {r['periodo']}", f"El avance acumulado ({fnum(acum)}) supera la meta anual ({fnum(meta)}).")
                cods.append("K7")
            if r["estado"] == "Cumplido" and p and a < p:
                H("INFO", "K8", f"{idk} · {r['periodo']}", f"Estado «Cumplido» con {fnum(a)} de {fnum(p)} programado en el mes.")
            r["validacion"] = ";".join(cods)
            filas.append(r)
        # meses faltantes y acumulado parcial
        pers = [r["periodo"] for r in rs]
        a0, m0 = int(pers[0][:4]), int(pers[0][5:])
        a1, m1 = int(pers[-1][:4]), int(pers[-1][5:])
        esperados = [f"{y}-{m:02d}" for y in range(a0, a1 + 1) for m in range(1, 13) if (y, m) >= (a0, m0) and (y, m) <= (a1, m1)]
        faltan = [p for p in esperados if p not in pers]
        if faltan:
            H("ADVERTENCIA", "K10", idk, "Meses sin reporte entre el primero y el último: " + ", ".join(faltan) + ".")
        if m0 != 1 and not es_pct(c):
            H("INFO", "K11", idk, f"El histórico empieza en {per_txt(pers[0])}: el avance acumulado {pers[0][:4]} solo suma desde ese mes "
              "(si hubo logros antes, falta cargarlos o registrar el acumulado inicial).")
        ultimo[idk] = rs[-1]

    # ---- escribir datos
    escribir_csv(os.path.join(dir_area, "catalogo.csv"),
                 [{**c, "meta_anual": fnum(c.get("meta_anual"))} for c in cat], CAMPOS_CAT)
    filas.sort(key=lambda r: (r["id_kpi"], r["periodo"]))
    escribir_csv(os.path.join(dir_area, "seguimiento.csv"), filas, CAMPOS_SEG)

    # ---- panorama destacado de kpis.html (mismo formato que ya usa la página)
    ruta_pan = os.path.join(args.datos, "kpis", cfg["panorama"])
    previas = leer_csv(ruta_pan)
    campos_pan = list(previas[0].keys()) if previas else []
    # Solo los KPI del formato de seguimiento: el área se reescribe completa cada mes
    # (Iniciativas Económicas y Resoluciones ya no forman parte de Pueblos y Nacionalidades).
    otras_areas = [r for r in previas if r.get("area") != cfg["area_kpis"]]
    pan = []
    for idk, c in ids.items():
        u = ultimo.get(idk)
        if not u:
            continue
        base = {k: "" for k in campos_pan}
        base.update({"area": cfg["area_kpis"], "registro_id": idk, "indicador": c["indicador"], "tipo_registro": "desagregacion_actual",
                     "icono": "doc", "color": "var(--azul)", "color_claro": "var(--azul-pale)",
                     "headline": f"{int(num(u['avance_acumulado'])):,}".replace(",", "."),
                     "headline_sufijo": f"logrado(s) a {per_txt(u['periodo'])} · meta anual {fnum(c['meta_anual'])}",
                     "unidad": "count", "tendencia_direccion": "up" if u["estado"] in ("Cumplido", "En proceso") else "down",
                     "tendencia_texto": f"Estado reportado: {u['estado'] or 'sin estado'} · {per_txt(u['periodo'])}",
                     "meta_pct": str(round(num(u["avance_anual_pct"]) or 0)), "meta_label": f"de la meta anual {c['anio_meta']}",
                     "tabulado_encabezado_categoria": f"Corte {per_txt(u['periodo'])}", "tabulado_encabezado_valor": "Valor",
                     "impacto_texto": c["resultado_esperado"], "es_dato_real": "TRUE",
                     "nota": f"Dato real: formato de seguimiento de KPI ({os.path.basename(args.archivo)})."})
        tab = [("Meta programada del mes", u["meta_programada"]), ("Meta alcanzada del mes", u["meta_alcanzada"]),
               ("Cumplimiento del mes", (u["cumplimiento_mensual_pct"].replace(".", ",") + " %") if u["cumplimiento_mensual_pct"] else "No aplica"),
               ("Avance acumulado", f"{u['avance_acumulado']} de {fnum(c['meta_anual'])}")]
        if es_pct(c):
            pc = lambda v: f"{float(v):.1f}".replace(".", ",") + " %" if v not in ("", None) else "No aplica"
            num_l, den_l = c.get("etiqueta_numerador") or "Numerador", c.get("etiqueta_denominador") or "Denominador"
            base.update({"headline": pc(u["cumplimiento_mensual_pct"]), "unidad": "percent",
                         "headline_sufijo": f"{num_l.lower()} sobre el total del mes · {per_txt(u['periodo'])}",
                         "tendencia_direccion": "", "tendencia_texto": f"{fnum(u['meta_alcanzada'])} de {fnum(u['meta_programada'])} {den_l.lower()}",
                         "meta_pct": "", "meta_label": ""})
            tab = [(num_l, u["meta_alcanzada"]), (den_l, u["meta_programada"]),
                   ("Resultado del mes", pc(u["cumplimiento_mensual_pct"])),
                   (f"Acumulado {u['periodo'][:4]}", pc(u["cumplimiento_fecha_pct"]))]
        for cat_, val in tab:
            pan.append({**base, "desagregacion_categoria": cat_, "desagregacion_valor": val})
    if campos_pan:
        escribir_csv(ruta_pan, otras_areas + pan, campos_pan)

    # ---- reporte
    ult = max((r["periodo"] for r in filas), default="sin-datos")
    os.makedirs(args.reportes, exist_ok=True)
    rep = os.path.join(args.reportes, f"validacion_{args.area}_{ult}.md")
    orden = {"ERROR": 0, "ADVERTENCIA": 1, "INFO": 2}
    n = defaultdict(int)
    for h in hall:
        n[h[0]] += 1
    L = [f"# Validación KPI · {cfg['nombre']} — último corte {ult}", "",
         f"Archivo: {os.path.basename(args.archivo)} · meses nuevos: {nuevos} · meses reemplazados: {reemplazados} · "
         f"filas en el histórico: {len(filas)}", "", f"**Errores: {n['ERROR']}** · Advertencias: {n['ADVERTENCIA']} · Informativos: {n['INFO']}", "",
         "## Hallazgos", "", "| Severidad | Código | Dónde | Detalle |", "|---|---|---|---|"]
    L += [f"| {s} | {c} | {d} | {t.replace('|', '/')} |" for s, c, d, t in sorted(hall, key=lambda h: (orden[h[0]], h[1]))] or ["| — | — | — | Sin hallazgos |"]
    L += ["", "## Último corte por KPI", "", "| KPI | Periodo | Programado | Logrado | % mes | Acumulado | % meta anual | Estado |", "|---|---|---:|---:|---:|---:|---:|---|"]
    for idk, u in ultimo.items():
        if es_pct(ids[idk]):
            L.append(f"| {idk} (porcentaje) | {u['periodo']} | {u['meta_programada']} (denominador) | {u['meta_alcanzada']} (numerador) | "
                     f"{u['cumplimiento_mensual_pct'] or '—'} | acumulado del año: {u['cumplimiento_fecha_pct'] or '—'} % | — | {u['estado'] or '—'} |")
        else:
            L.append(f"| {idk} | {u['periodo']} | {u['meta_programada']} | {u['meta_alcanzada']} | {u['cumplimiento_mensual_pct'] or '—'} | "
                     f"{u['avance_acumulado']} / {fnum(ids[idk]['meta_anual'])} | {u['avance_anual_pct']} % | {u['estado']} |")
    with open(rep, "w", encoding="utf-8") as fh:
        fh.write("\n".join(L) + "\n")
    print(f"{cfg['nombre']}: {len(cat)} KPI · {nuevos} mes(es) nuevo(s), {reemplazados} reemplazado(s) · histórico {len(filas)} filas. "
          f"Errores {n['ERROR']} · Advertencias {n['ADVERTENCIA']} · Info {n['INFO']}")
    print("Reporte:", rep)


if __name__ == "__main__":
    main()
